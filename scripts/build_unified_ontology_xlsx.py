"""통합 온톨로지(PHR+장치·환경+OCR·대화) → Excel(.xlsx)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E79"); ALT = PatternFill("solid", fgColor="EDF2F7")
SENS = PatternFill("solid", fgColor="F7E2D2"); DEV = PatternFill("solid", fgColor="EAF1E6")
thin = Side(style="thin", color="CCCCCC"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def sh(name, headers, rows, widths, sens_col=None):
    ws = wb.create_sheet(name); ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10); c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDER
    for i, row in enumerate(rows):
        ws.append(row)
        rs = "".join(str(x) for x in row)
        fill = SENS if (sens_col is not None and "민감" in str(row[sens_col] if len(row) > sens_col else "")) \
            else (DEV if ("디바이스" in rs or "환경" in rs or "OCR" in rs or "대화" in rs) and name == "필드분류(전체)" else (ALT if i % 2 else None))
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT, size=10); c.alignment = Alignment(vertical="center", wrap_text=True); c.border = BORDER
            if fill: c.fill = fill
    for col, w in zip("ABCDEFG", widths): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; return ws

wb.remove(wb.active)
ws0 = wb.create_sheet("개요"); ws0["A1"] = "통합 데이터 온톨로지 (PHR·장치·환경·OCR·대화)"
ws0["A1"].font = Font(name=FONT, bold=True, size=15, color="1F4E79")
for i, n in enumerate([
    "", "정본: docs/ontology/*.ttl (phr·device·context) — 통합 741 트리플·120 노드.", "작성: 2026-06-21",
    "", "핵심: 모든 소스가 signalKey로 같은 참조범위를 조인 → 해석·안전 코어는 하나.",
    "활용: 밴드 / deny / 시간산술 / 이력ack / 내부 / 웰니스(기저선) | 민감 행=주황, 디바이스·환경·OCR·대화 행=연두."], start=2):
    ws0[f"A{i}"] = n; ws0[f"A{i}"].font = Font(name=FONT, size=10)
ws0.column_dimensions["A"].width = 115

sh("클래스", ["클래스", "상위", "라벨", "설명"], [
    ["phr:Person/PHR", "—", "사람·개인건강기록", "주체"],
    ["phr:HealthRecord", "—", "건강기록(추상)", "PHR 8 카테고리 상위"],
    ["  ⊃ 8종", "HealthRecord", "검진·질병·가족·생활·암검진·리포트·처방·진료", ""],
    ["phr:Observation", "—", "관찰값(추상)", "검진 측정값 상위"],
    ["phr:Device", "—", "장치(추상)", "혈압기·체중계·워치·A1·공기질"],
    ["phr:DeviceMeasurement", "Observation", "장치 측정값", "해석 코어 재사용(⊑Observation)"],
    ["  ⊃ WatchSignal", "DeviceMeasurement", "워치 신호(웰니스)", "기저선 대비·임상밴드 금지"],
    ["phr:EnvironmentReading", "—", "실내 환경 측정", "공기질(PM·CO2 등) — 신체 아님"],
    ["phr:OCRDocument", "—", "OCR 문서(입력 양식)", "처방전·검사지·진단서·약봉투"],
    ["phr:Conversation", "—", "대화(LLM 세션)", "턴·의도·수집맥락"],
    ["phr:DerivedContext", "—", "수집 맥락(축)", "증상·위험신호·환자맥락"],
    ["phr:ReferenceRange/BandLabel/Trend", "—", "참조범위·밴드·추세", "공유 해석 코어"],
    ["phr:WellnessLabel/PersonalBaseline", "—", "웰니스 라벨·기저선", "웰니스 신호 해석"],
    ["phr:CrossSignalCombo", "—", "교차신호 화이트리스트", "출처 첨부 조합만(I9)"],
], [26, 18, 34, 40])

sh("필드분류(전체)", ["대분류", "필드", "데이터타입", "단위", "민감도", "활용"], [
    # PHR 검진
    ["PHR 검진·계측", "신장·체중·허리둘레·BMI", "수치", "cm/kg", "일반", "밴드(추세)"],
    ["PHR 검진·혈압혈액", "혈압·혈색소·공복혈당·총콜레스테롤·HDL·중성지방·크레아티닌", "수치", "—", "일반", "밴드"],
    ["PHR 검진·혈압혈액", "LDL · eGFR", "수치", "—", "일반", "deny(맥락 의존)"],
    ["PHR 검진·간요영상골", "AST·ALT·γ-GTP", "수치", "U/L", "일반", "밴드(검사실)"],
    ["PHR 검진·간요영상골", "요단백 · 골밀도 T-score", "정성/점수", "—", "일반", "deny"],
    ["PHR 검진·간요영상골", "흉부X선 · 종합판정", "판정문/등급", "—", "일반", "ack"],
    ["PHR 질병이력", "고혈압·당뇨·이상지질·뇌졸중·심장·COPD", "불리언+시기", "—", "일반", "이력ack·shaping"],
    ["PHR 질병이력", "결핵·B/C형간염 · 우울·정신질환", "불리언", "—", "민감", "내부(표면화0)"],
    ["PHR 가족력", "뇌졸중·심장·고혈압·당뇨 / 암", "불리언+관계", "—", "일반/민감", "위험요인 / direct만"],
    ["PHR 생활습관", "흡연·음주·신체활동", "범주/수치", "갑년", "일반", "위험요인·교차"],
    ["PHR 암검진", "위·대장·간·폐·유방·자궁: 수검일·결과", "날짜/판정", "—", "일반/민감", "시간산술(갭)"],
    ["PHR 처방", "약물계열·처방일수·내원유형", "범주/수치", "일", "계열따라 민감", "계열만·시간산술(리필)"],
    ["PHR 진료", "기관·진료일·진료과·입원/외래/응급·상병코드", "범주/날짜/코드", "—", "일반/민감", "시간산술·진료과·내부"],
    # 디바이스
    ["디바이스·혈압기/A1", "수축기·이완기·맥박·SpO2·체온", "수치", "mmHg 등", "일반", "밴드(의료근접)"],
    ["디바이스·A1", "스트레스 지수", "수치", "—", "일반", "내부(참조범위 없음)"],
    ["디바이스·체중계", "체중·BMI", "수치", "kg", "일반", "밴드·추세"],
    ["디바이스·체중계", "체지방률·근육량·내장지방", "수치", "%/kg", "일반", "밴드(추정·천장낮음)"],
    ["디바이스·워치(웰니스)", "안정시심박·수면·활동·걸음·운동·호흡수", "수치", "—", "일반", "웰니스(기저선)"],
    ["디바이스·워치(웰니스)", "HRV", "수치", "ms", "일반", "웰니스·교차 제외"],
    ["디바이스·워치(웰니스)", "워치 SpO2 · 심전도/불규칙맥 · 합성점수", "—", "—", "일반", "deny(임상판정 금지)"],
    # 환경
    ["환경·공기질", "종합·PM10·PM2.5·PM1·NOx·VOC·HCHO·CO2", "수치/등급", "µg/㎥·ppm", "일반", "환경 교차 참고·환기"],
    ["환경·공기질", "실내 온도·습도", "수치", "℃/%", "일반", "참고(서술)"],
    ["환경·공기청정기", "필터수명·운전모드·측정시각", "수치/범주", "—", "일반", "환기 안내 보조"],
    # OCR
    ["OCR(입력 양식)", "처방전·검사지·진단서·약봉투", "이미지→추출", "—", "원천 따름", "확인 후 PHR 합류"],
    ["OCR(입력 양식)", "추출 필드(필드별 신뢰도)", "값+confidence", "—", "—", "저신뢰·단위실패=차단"],
    # 대화
    ["대화(LLM)", "증상프로파일·위험신호·환자맥락(나이·성별·복약·생활)", "텍스트→축", "—", "민감(건강텍스트)", "shaping(진단 아님)"],
    ["대화(LLM)", "의도(응급/증상/정보/바이탈) · 진술수치", "범주/수치", "—", "—", "진술수치=밴드 미활성(I4)"],
], [18, 38, 14, 12, 16, 22], sens_col=4)

sh("객체속성", ["속성", "Domain", "Range", "설명"], [
    ["hasObservation", "HealthCheckup", "Observation", "검진→관찰값"],
    ["interpretedAs", "Observation", "BandLabel", "관찰값→밴드(★, A3·A5)"],
    ["partOfTrend / comparedTo", "Observation", "Trend·ReferenceRange", "추세·참조범위"],
    ["measuredBy / hasDeviceGrade", "DeviceMeasurement", "Device·DeviceGrade", "측정 장치·등급"],
    ["interpretedAsWellness / comparedToBaseline", "WatchSignal", "WellnessLabel·PersonalBaseline", "웰니스 해석(B1)"],
    ["hasEnvGrade", "EnvironmentReading", "EnvGrade", "공기질 등급"],
    ["partOfCombo", "DeviceMeasurement", "CrossSignalCombo", "교차 조합(I9)"],
    ["extractedFrom / feedsInto", "ExtractedField·OCRDocument", "OCRDocument·PHR레코드", "OCR→PHR 합류"],
    ["hasIntent / collectedAxis", "UserUtterance·Conversation", "Intent·DerivedContext", "의도·수집맥락"],
    ["records / diagnosis", "이력·진료", "Condition", "질환 참조"],
    ["hasSensitivity / hasProvenance", "(횡단)", "Sensitivity·Provenance", "민감도·출처"],
], [24, 24, 22, 30])

sh("데이터속성", ["속성", "Domain", "타입", "설명"], [
    ["value", "Observation", "decimal", "원시값 — 내부전용·외부 미전송(A5/B7)"],
    ["clinicalLabel / labelUser", "Observation", "string/enum", "임상라벨(내부) / 노출 중립라벨"],
    ["signalKey", "측정 클래스", "string", "참조범위 조인 키(공유)"],
    ["deviceGrade", "Device", "enum", "의료근접/소비자/웰니스/환경"],
    ["fieldConfidence / userConfirmed", "ExtractedField/OCR", "decimal/bool", "OCR 필드 신뢰·확인(C1·C3)"],
    ["daysSupplied / ageAtOnset / packYears", "처방·가족력·생활", "int/decimal", "리필·발병연령·갑년"],
], [22, 22, 14, 38])

sh("통제어휘", ["개념", "허용값"], [
    ["BandLabel(노출)", "안정 · 주의 · 경고"],
    ["WellnessLabel", "평소대비높음 · 평소대비낮음 · 평소수준 · 참고 (본인 기저선)"],
    ["Trend", "안정유지 · 지속상승 · 지속저하 · 불안정반복 (개선/악화 제외)"],
    ["Sensitivity", "일반 · 민감(정신·암·감염병)"],
    ["Utilization", "밴드 · deny · 시간산술 · 이력ack · 내부 · 웰니스"],
    ["DeviceGrade", "의료근접 · 소비자 · 웰니스 · 환경"],
    ["EnvGrade", "좋음 · 보통 · 나쁨 · 매우나쁨"],
    ["Intent(대화)", "응급 · 증상 · 정보 · 바이탈 · 비핵심"],
    ["ConfirmState(OCR)", "추출됨 · 확인필요 · 차단 · 사용자확인"],
    ["VisitType / ScreeningResult / Provenance", "외래·입원·응급 / 정상·이상소견·유보 / 공단·심평원·파트너·OCR"],
], [22, 74])

sh("공리", ["#", "안전 규칙"], [
    ["A1", "PHR: 민감 질환(정신·암·감염병) 표면화 금지 — 내부 신호만"],
    ["A2", "PHR: deny(eGFR·골밀도·요단백·LDL) 밴드 금지 → '수치는 의료진과'"],
    ["A3", "노출 라벨 ⊆ {안정,주의,경고}; 임상/질환명 라벨 내부 전용"],
    ["A4", "추세는 데이터 패턴만 — 개선/악화(건강판단) 금지"],
    ["A5", "원시 측정값 내부 전용 — 외부(LLM/국외) 미전송"],
    ["A6", "암검진 갭 = 오늘 − 마지막수검일 − 권장주기"],
    ["A7", "처방 노출 = 약물계열 ∧ 과거시제 ∧ (정신과·항감염 계열 내부)"],
    ["A8", "표면화 = 민감도 일반 ∧ 질의 관련 ∧ 신선일 때만"],
    ["B1", "장치: deviceGrade=웰니스 ⇒ 임상밴드 금지, 본인 기저선 WellnessLabel만"],
    ["B2", "장치: ECG·불규칙맥·워치SpO2(저산소)·낙상·합성점수 deny(기기 알림 패스스루만)"],
    ["B3", "장치: 워치 웰니스 신호(HRV·스트레스·수면) 교차 입력 영구 제외(유사과학 차단)"],
    ["B4", "장치: 교차 조합 = 화이트리스트+출처 필수, 미등재 생성 금지"],
    ["B5", "환경: 진단 아님 — '환경 교차 참고+환기 권유' 천장, 실내 한정"],
    ["B6", "장치: 사용자 채팅 타이핑 수치 개인화 미활성(연동만)"],
    ["B7", "장치: 원시값·고해상 파형(ECG·RR·hypnogram) 저장·전송 금지(집계만)"],
    ["B8", "장치: 출처충돌 → 신뢰위계 + 재측정 권유"],
    ["B9", "장치: 체지방·근육·내장지방(추정) 추세 위주·천장 낮음"],
    ["C1", "OCR: 사용자 확인 전 활용 차단(확인 게이트)"],
    ["C2", "OCR: 진단서→질병이력 자동생성 금지(사람 확인분만)"],
    ["C3", "OCR: 민감질환·수치·단위·약물용량 confidence 무관 강제 재확인"],
    ["C4", "OCR: 단위 추출 실패/불일치 ⇒ 라벨 금지(blocked)"],
    ["C5", "OCR: user_confirmed 후 디바이스급 승격(신뢰위계)"],
    ["D1", "대화: 사용자 진술 수치 밴드 미활성(증상 맥락으로만, I4)"],
    ["D2", "대화: 매 턴 응급 재평가 불변(무상태)"],
    ["D3", "대화: 수집 맥락은 답변 shaping 입력 — 진단 단정 아님"],
    ["D4", "대화: 멀티턴 — 이미 말한 축 재질문 금지, 새 응급은 재평가"],
    ["D5", "대화: 자유텍스트(건강) 동의·마스킹, 원문 우회 금지"],
], [6, 90])

import os
os.makedirs("docs/ontology", exist_ok=True)
wb.save("docs/ontology/통합_온톨로지_2026-06-21.xlsx")
print("saved xlsx")
