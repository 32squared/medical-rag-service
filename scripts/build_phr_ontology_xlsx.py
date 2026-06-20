"""PHR 온톨로지 → Excel(.xlsx). 다중 시트 참조 카탈로그(수식 없음)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
ALT_FILL = PatternFill("solid", fgColor="EDF2F7")
SENS_FILL = PatternFill("solid", fgColor="F7E2D2")
thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def sheet(name, headers, rows, widths, sens_col=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, row in enumerate(rows):
        ws.append(row)
        is_sens = sens_col is not None and len(row) > sens_col and "민감" in str(row[sens_col])
        for c in ws[ws.max_row]:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = BORDER
            if is_sens:
                c.fill = SENS_FILL
            elif i % 2:
                c.fill = ALT_FILL
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return ws

wb.remove(wb.active)

# 1. 표지/개요
ws0 = wb.create_sheet("개요")
ws0["A1"] = "PHR 데이터 온톨로지"
ws0["A1"].font = Font(name=FONT, bold=True, size=16, color="1F4E79")
notes = [
    "", "공단·심평원 건강검진/진료/처방 데이터를 진단 없이 안전하게 활용하기 위한 개념 모델.",
    "작성: 2026-06-21 · 정본: docs/ontology/phr-ontology.ttl",
    "",
    "시트 구성:",
    " · 클래스 — 개념 계층(taxonomy)",
    " · 필드분류(전체) — PHR 전 필드 × 타입·민감도·활용 (핵심 카탈로그)",
    " · 객체속성 — 개념 간 관계",
    " · 데이터속성 — 속성(attribute)",
    " · 통제어휘 — 허용값(enum)",
    " · 공리 — 안전 규칙(A1~A8)",
    "",
    "활용 표기: 밴드=구간라벨 / deny=밴드금지·의료진위임 / 시간산술=시기계산 / ack=이력표시 / 내부=비표시·내부신호",
    "민감도: 일반 / 민감(정신·암·감염병) — 민감 행은 주황 음영.",
]
for i, n in enumerate(notes, start=2):
    ws0[f"A{i}"] = n
    ws0[f"A{i}"].font = Font(name=FONT, size=10, bold=(n.endswith("구성:")))
ws0.column_dimensions["A"].width = 110

# 2. 클래스 계층
sheet("클래스",
    ["클래스", "상위(부모)", "라벨", "설명"],
    [
        ["phr:Person", "—", "사람", "기록의 주체"],
        ["phr:PHR", "—", "개인건강기록", "한 사람의 전체 기록 집합"],
        ["phr:HealthRecord", "—", "건강기록(추상)", "8개 기록 카테고리의 상위"],
        ["phr:HealthCheckup", "HealthRecord", "건강검진", "검진 이벤트(관찰값 다수 포함)"],
        ["phr:DiseaseHistory", "HealthRecord", "질병이력", "과거 진단 이력"],
        ["phr:FamilyHistory", "HealthRecord", "가족력", "직계 가족 질환 이력"],
        ["phr:Lifestyle", "HealthRecord", "생활습관", "흡연·음주·신체활동"],
        ["phr:CancerScreening", "HealthRecord", "암검진", "6대 암 수검 이벤트"],
        ["phr:HealthReport", "HealthRecord", "건강리포트", "종합등급·건강나이 등 파생"],
        ["phr:Prescription", "HealthRecord", "처방", "약물계열·처방일수"],
        ["phr:TreatmentVisit", "HealthRecord", "진료이력", "요양기관·진료과·상병"],
        ["phr:Observation", "—", "관찰값(추상)", "검진 측정값의 상위"],
        ["phr:Anthropometric", "Observation", "계측", "신장·체중·허리·BMI"],
        ["phr:VitalMeasure", "Observation", "활력", "혈압"],
        ["phr:BloodTest", "Observation", "혈액검사", "혈색소·혈당·지질·신장"],
        ["phr:LiverFunctionTest", "Observation", "간기능검사", "AST·ALT·GGT"],
        ["phr:UrineTest", "Observation", "요검사", "요단백"],
        ["phr:ImagingResult", "Observation", "영상검사", "흉부X선"],
        ["phr:BoneDensity", "Observation", "골밀도", "T-score"],
        ["phr:OverallJudgment", "Observation", "종합판정", "검진 종합소견·등급"],
        ["phr:Condition", "—", "질환/상병", "질병/가족력/진료가 참조"],
        ["phr:DrugClass", "—", "약물계열(ATC)", "약물명 대신 계열로 추상화"],
        ["phr:ReferenceRange", "—", "참조범위", "밴드 정의(해석 입력)"],
        ["phr:BandLabel", "—", "밴드 라벨", "안정/주의/경고(노출)"],
        ["phr:Trend", "—", "추세", "데이터 패턴(개선/악화 제외)"],
    ],
    [26, 18, 16, 46])

# 3. 필드분류 (전체 카탈로그) — 핵심
sheet("필드분류(전체)",
    ["대분류", "필드", "데이터타입", "단위", "민감도", "활용"],
    [
        ["건강검진·계측", "신장(키)", "수치", "cm", "일반", "BMI 입력(내부)"],
        ["건강검진·계측", "체중", "수치", "kg", "일반", "추세·BMI"],
        ["건강검진·계측", "허리둘레", "수치", "cm", "일반", "밴드(성별)"],
        ["건강검진·계측", "체질량지수(BMI)", "수치(파생)", "kg/m²", "일반", "밴드"],
        ["건강검진·계측", "시력·청력", "범주", "—", "일반", "ack"],
        ["건강검진·혈압혈액", "수축기혈압", "수치", "mmHg", "일반", "밴드"],
        ["건강검진·혈압혈액", "이완기혈압", "수치", "mmHg", "일반", "밴드"],
        ["건강검진·혈압혈액", "혈색소(Hb)", "수치", "g/dL", "일반", "밴드(성별)"],
        ["건강검진·혈압혈액", "공복혈당", "수치", "mg/dL", "일반", "밴드"],
        ["건강검진·혈압혈액", "총콜레스테롤", "수치", "mg/dL", "일반", "밴드"],
        ["건강검진·혈압혈액", "HDL콜레스테롤", "수치", "mg/dL", "일반", "밴드(방향반전)"],
        ["건강검진·혈압혈액", "LDL콜레스테롤", "수치", "mg/dL", "일반", "deny(위험계층 가변)"],
        ["건강검진·혈압혈액", "중성지방", "수치", "mg/dL", "일반", "밴드(공복)"],
        ["건강검진·혈압혈액", "크레아티닌", "수치", "mg/dL", "일반", "밴드(성별)"],
        ["건강검진·혈압혈액", "eGFR", "수치", "mL/min/1.73㎡", "일반", "deny(급성/만성·연령)"],
        ["건강검진·간요영상골", "AST", "수치", "U/L", "일반", "밴드(검사실ULN)"],
        ["건강검진·간요영상골", "ALT", "수치", "U/L", "일반", "밴드"],
        ["건강검진·간요영상골", "γ-GTP", "수치", "U/L", "일반", "밴드(성별)"],
        ["건강검진·간요영상골", "요단백", "정성", "음성/±/1+~", "일반", "deny(일시적 흔함)"],
        ["건강검진·간요영상골", "흉부X선 판정", "판정문", "—", "소견따라", "ack·내부"],
        ["건강검진·간요영상골", "골밀도 T-score", "수치(점수)", "SD", "일반", "deny(부위·연령·성별)"],
        ["건강검진·간요영상골", "구강(치아우식 등)", "범주", "—", "일반", "ack"],
        ["건강검진·간요영상골", "종합판정·소견", "등급/판정문", "—", "일반", "ack(등급)"],
        ["질병이력", "고혈압·당뇨·이상지질·뇌졸중·심장병·COPD", "불리언+시기", "—", "일반", "이력ack·shaping"],
        ["질병이력", "결핵·B형/C형간염", "불리언", "—", "민감(감염병)", "내부(표면화0)"],
        ["질병이력", "우울증·정신질환", "불리언", "—", "민감(정신)", "내부(표면화0)"],
        ["가족력", "뇌졸중·심장병·고혈압·당뇨", "불리언+관계", "—", "일반", "위험요인 묶음"],
        ["가족력", "암", "불리언+관계", "—", "민감", "direct+동의 시만"],
        ["생활습관", "흡연(상태·흡연량·갑년)", "범주/수치", "갑년", "일반", "위험요인·폐검진 교차"],
        ["생활습관", "음주(빈도·1회량)", "범주/수치", "—", "일반", "간수치 교차"],
        ["생활습관", "신체활동(빈도·강도)", "범주", "—", "일반", "활동 안내"],
        ["암검진", "위·대장·간·폐·유방·자궁: 수검여부·수검일", "불리언·날짜", "—", "일반", "시간산술(권장주기 갭)"],
        ["암검진", "검진 결과(정상/이상소견/유보)", "판정", "—", "민감 가능", "표면화 주의·내부"],
        ["건강리포트", "종합등급·건강나이·신체점수·생활습관점수·위험요인", "등급/수치", "—", "일반", "ack·요약카드"],
        ["처방", "약물 성분/계열(ATC)", "범주", "—", "계열따라 민감(정신·항감염)", "계열만·과거시제"],
        ["처방", "처방일수·투약일수", "수치", "일", "일반", "시간산술(리필)"],
        ["처방", "처방기관·조제일", "텍스트·날짜", "—", "일반", "내부"],
        ["진료", "요양기관(종별)", "범주", "—", "일반", "기관 클래스만"],
        ["진료", "진료일", "날짜", "—", "일반", "시간산술(경과)"],
        ["진료", "진료과목", "범주", "—", "일반", "진료과 추적"],
        ["진료", "입원/외래/응급 구분", "범주", "—", "일반", "빈도(내부)"],
        ["진료", "상병(진단)코드", "코드(ICD)", "—", "민감 가능", "내부·민감분류 입력"],
        ["진료", "본인부담금", "수치", "원", "일반", "(미활용)"],
    ],
    [16, 34, 14, 14, 18, 22], sens_col=4)

# 4. 객체속성
sheet("객체속성",
    ["속성(Property)", "Domain(주어)", "Range(목적어)", "설명"],
    [
        ["hasPHR", "Person", "PHR", "사람→개인건강기록"],
        ["hasRecord", "PHR", "HealthRecord", "PHR→기록(하위: hasCheckup 등)"],
        ["hasObservation", "HealthCheckup", "Observation", "검진→관찰값"],
        ["hasJudgment", "HealthCheckup", "OverallJudgment", "검진→종합판정"],
        ["comparedTo", "Observation", "ReferenceRange", "관찰값→참조범위(해석 입력)"],
        ["interpretedAs", "Observation", "BandLabel", "관찰값→밴드 라벨 (★핵심, A3/A5)"],
        ["partOfTrend", "Observation", "Trend", "관찰값→추세"],
        ["definesBand", "ReferenceRange", "Band", "참조범위→밴드 구간"],
        ["hasDrugClass", "Prescription", "DrugClass", "처방→약물계열"],
        ["prescribedAt", "Prescription", "Institution", "처방→기관"],
        ["records", "HealthRecord", "Condition", "질병/가족력/진료→질환"],
        ["screensFor", "CancerScreening", "Condition", "암검진→대상 암"],
        ["recommendedCycle", "CancerScreening", "RecommendedCycle", "암검진→권장주기"],
        ["atInstitution", "TreatmentVisit", "Institution", "진료→기관"],
        ["inDepartment", "TreatmentVisit", "MedicalDepartment", "진료→진료과"],
        ["diagnosis", "TreatmentVisit", "Condition", "진료→상병"],
        ["hasSensitivity", "(횡단)", "Sensitivity", "민감도 분류"],
        ["hasProvenance", "(횡단)", "Provenance", "데이터 출처"],
    ],
    [20, 18, 20, 40])

# 5. 데이터속성
sheet("데이터속성",
    ["속성", "Domain", "타입", "설명"],
    [
        ["value", "Observation", "decimal", "원시 측정값 — 내부전용·외부 미전송(A5)"],
        ["qualitativeValue", "UrineTest 등", "string", "정성 결과(음성/1+ 등)"],
        ["observedAt", "(관측 클래스)", "date", "관측·측정 시점"],
        ["daysSupplied", "Prescription", "integer", "처방일수(리필 산술)"],
        ["isPresent", "DiseaseHistory/FamilyHistory", "boolean", "이력 유무"],
        ["ageAtOnset", "FamilyHistory", "integer", "가족 발병 연령"],
        ["packYears", "Lifestyle", "decimal", "흡연 갑년"],
        ["clinicalLabel", "Observation", "string", "임상/질환명 라벨 — 내부전용(A3)"],
        ["labelUser", "Observation", "enum", "노출 라벨(안정/주의/경고)"],
    ],
    [18, 26, 12, 44])

# 6. 통제어휘
sheet("통제어휘",
    ["개념", "허용값"],
    [
        ["BandLabel(노출)", "안정 · 주의 · 경고"],
        ["Trend", "안정유지 · 지속상승 · 지속저하 · 불안정반복  (개선/악화 제외)"],
        ["Sensitivity", "일반 · 민감(정신·암·감염병)"],
        ["Utilization", "밴드 · deny · 시간산술 · 이력ack · 내부"],
        ["VisitType", "외래 · 입원 · 응급"],
        ["ScreeningResult", "정상 · 이상소견 · 판정유보"],
        ["Provenance", "공단 · 심평원 · 파트너앱 · OCR"],
        ["Sex / Context", "남·여 / home·clinic"],
    ],
    [22, 72])

# 7. 공리
sheet("공리",
    ["#", "안전 규칙(공리)"],
    [
        ["A1", "민감 질환(정신·암·감염병)은 표면화 금지 — 내부 신호만"],
        ["A2", "deny 목록(eGFR·골밀도·요단백·LDL)은 밴드 라벨 금지 → '수치는 의료진과'"],
        ["A3", "노출 라벨 ⊆ {안정,주의,경고}; 임상/질환명 라벨은 내부 전용"],
        ["A4", "추세는 데이터 패턴만 — '개선/악화'(건강 판단) 금지"],
        ["A5", "원시 측정값은 내부 전용 — 외부(LLM/국외) 미전송"],
        ["A6", "암검진 갭 = 오늘 − 마지막수검일 − 권장주기 (시간산술)"],
        ["A7", "처방 노출 = 약물계열 ∧ 과거시제 ∧ (정신과·항감염 계열은 내부)"],
        ["A8", "표면화 = 민감도 일반 ∧ 질의 관련 ∧ 신선(fresh)일 때만"],
    ],
    [6, 88])

import os
os.makedirs("docs/ontology", exist_ok=True)
wb.save("docs/ontology/PHR_온톨로지_2026-06-21.xlsx")
print("saved")
