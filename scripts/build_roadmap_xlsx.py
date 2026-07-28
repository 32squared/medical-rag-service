# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
FONT="Arial"; HEAD=PatternFill("solid",fgColor="1F4E79")
QF={"Q1":"DCEAF5","Q2":"E3F0E0","Q3":"FBEFD9","Q4":"F2E5F0"}
thin=Side(style="thin",color="CCCCCC"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); ws=wb.active; ws.title="출시 로드맵"
ws.append(["분기","서비스","층","근거","선결조건"])
for c in ws[1]:
    c.font=Font(name=FONT,bold=True,color="FFFFFF",size=10); c.fill=HEAD
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
rows=[
 ["Q1","1.1 혈압 안내","1","엔진 완성·고가치","디바이스 연동 결정"],
 ["Q1","2.1 검진 결과 풀이","2","킬러·밴드 엔진 일부","검진밴드 시드·PHR 출처"],
 ["Q1","2.3 검진·복약·진료 알리미","2","저리스크·시간산술","PHR 출처"],
 ["Q1","1.4 공기질·환기","1","환경 시드 보유","공기청정기 연동"],
 ["Q2","1.2 체중·체성분 추세","1","추세 엔진 완성","체중계 연동·이력저장"],
 ["Q2","2.2 연도별 추세","2","변화 인식","다년 검진 데이터"],
 ["Q2","2.5 내 건강 요약 카드","2","제품 허브","PHR 통합·UI"],
 ["Q2","3.1 증상+최근 측정","3","결합 상담 시작 ★","측정 연동"],
 ["Q2","3.5 병원 갈 때 준비","3","실용 차별","PHR+대화"],
 ["Q3","3.2 증상+검진 이력","3","개인화 상담","민감 분류"],
 ["Q3","3.4 증상+복약 이력","3","복약 안전","약물계열 정규화"],
 ["Q3","1.3 컨디션 추세(워치)","1","웰니스","워치 연동·기저선·정확도 고지"],
 ["Q3","4.2 대사 관리 묶음","4","위험요인 통합","여러 소스 연동"],
 ["Q3","5.1 호흡기·알레르기 케어","5","환경×이력","공기질×질병이력"],
 ["Q4","4.1 컨디션 저하 조기신호","4","워치 교차(공인 조합)","워치 신호·화이트리스트"],
 ["Q4","4.3 활동·수면 코칭","4","생활 코칭","워치 집계"],
 ["Q4","3.3 증상+환경","3","환경 결합","공기질 연동"],
 ["Q4","5.2 수면·집중 환경","5","환기 안내","CO2 연동"],
 ["Q4","6.1 위험요인 묶음","6","예방 관리","검진+가족력+측정"],
 ["Q4","6.2 암검진 갭(위험군)","6","위험군 검진","질병이력 교차"],
 ["Q4","2.4 종이기록 디지털화(OCR)","2","입력 확장","OCR 확인게이트"],
]
for i,r in enumerate(rows):
    ws.append(r)
    for c in ws[ws.max_row]:
        c.font=Font(name=FONT,size=10); c.alignment=Alignment(vertical="center",wrap_text=True); c.border=B
        c.fill=PatternFill("solid",fgColor=QF[r[0]])
for col,w in zip("ABCDE",[8,30,7,26,30]): ws.column_dimensions[col].width=w
ws.freeze_panes="A2"
ws2=wb.create_sheet("분기 테마")
for i,(q,th) in enumerate([("Q1","저리스크·고가치 단일 — 혈압·검진풀이·알리미·환기 (연동 결정 직후)"),
  ("Q2","추세·요약 + 결합상담 시작 — 체중추세·연도추세·요약카드·증상+측정·병원준비"),
  ("Q3","개인화 심화 + 워치 — 증상+검진/복약·컨디션추세·대사묶음·호흡기환경"),
  ("Q4","교차신호·환경·예방·OCR — 조기신호·활동코칭·증상+환경·위험요인·암검진갭·OCR")],start=1):
    ws2[f"A{i}"]=q; ws2[f"B{i}"]=th
    ws2[f"A{i}"].font=Font(name=FONT,bold=True,size=11); ws2[f"B{i}"].font=Font(name=FONT,size=10)
ws2.column_dimensions["A"].width=8; ws2.column_dimensions["B"].width=92
wb.save("docs/서비스_출시로드맵_2026-06-21.xlsx"); print("saved roadmap", ws.max_row-1)
