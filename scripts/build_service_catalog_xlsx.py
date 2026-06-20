"""통합 온톨로지 기반 서비스 카탈로그 → Excel."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
FONT="Arial"; HEAD=PatternFill("solid",fgColor="1F4E79"); ALT=PatternFill("solid",fgColor="EDF2F7")
COMBO=PatternFill("solid",fgColor="EAF1E6"); STAR=PatternFill("solid",fgColor="FCE8D5")
thin=Side(style="thin",color="CCCCCC"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); ws=wb.active; ws.title="서비스 카탈로그"
cols=["층/테마","서비스","유형","사용 데이터(소스)","사용자 가치","우선순위","안전(공리)"]
ws.append(cols)
for c in ws[1]:
    c.font=Font(name=FONT,bold=True,color="FFFFFF",size=10); c.fill=HEAD
    c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=B
rows=[
 ["1. 측정값 즉시해석","1.1 혈압 안내","단일","혈압기·A1","내 혈압 구간+상담 권유","높음","A2·밴드"],
 ["1. 측정값 즉시해석","1.2 체중·체성분 추세","단일","체중계","체중·BMI 흐름·급변 알림","높음","B9"],
 ["1. 측정값 즉시해석","1.3 컨디션 추세(웰니스)","단일","워치","수면·활동·심박 본인기준 추세","중","B1·B3"],
 ["1. 측정값 즉시해석","1.4 실내 공기질·환기","단일","공기질·공기청정기","환기 안내(진단 아님)","중","B5"],
 ["1. 측정값 즉시해석","1.5 다항목 바이탈","단일","A1 통합측정기","여러 항목 한번에 안내","중","밴드"],
 ["2. 기록 이해·챙김","2.1 검진 결과 풀이","단일(PHR)","PHR 검진","검진지 쉬운 해석 ★","★ 최고","A1·A2·A3"],
 ["2. 기록 이해·챙김","2.2 연도별 추세","단일(PHR)","PHR 검진(다년)","변화 인식·동기","높음","A4"],
 ["2. 기록 이해·챙김","2.3 검진·복약·진료 알리미","단일(PHR)","PHR 시간산술","놓치는 검진·재처방 챙김","★","A6"],
 ["2. 기록 이해·챙김","2.4 종이기록 디지털화","입력(OCR)","OCR","처방전·검사지→내 기록","중","C1~C5"],
 ["2. 기록 이해·챙김","2.5 내 건강 요약 카드","단일(PHR)","PHR 전반","흩어진 기록 한눈에","높음","A1·A3"],
 ["3. 증상+데이터 결합","3.1 증상+최근 측정","결합","대화×혈압기","내 측정 맥락 결합 상담 ★","★","D1·D3"],
 ["3. 증상+데이터 결합","3.2 증상+검진 이력","결합","대화×PHR","과거 이력 맥락 안내","높음","A1·D3"],
 ["3. 증상+데이터 결합","3.3 증상+환경","결합","대화×공기질","환기+진료 안내","중","B5·D3"],
 ["3. 증상+데이터 결합","3.4 증상+복약 이력","결합","대화×처방","약사 상담 연결","중","A7·D3"],
 ["3. 증상+데이터 결합","3.5 병원 갈 때 준비","결합","대화×PHR","진료 효율·자신감","높음","A3"],
 ["4. 교차신호 조기신호","4.1 컨디션 저하 조기신호","결합","워치(심박+수면+호흡)","감염 초기 패턴(공인 조합만)","중","B4"],
 ["4. 교차신호 조기신호","4.2 대사 관리 묶음","결합","가정혈압+체중+검진혈당","함께 볼 위험요인","높음","B4"],
 ["4. 교차신호 조기신호","4.3 활동·수면 코칭","결합","워치(활동+수면)","생활습관 일반 안내","중","B3·B4"],
 ["5. 환경×건강","5.1 호흡기·알레르기 케어","결합","공기질×질병이력","환기·외출 안내","중","B5·A1"],
 ["5. 환경×건강","5.2 수면·집중 환경","결합","CO2×대화","환기 권유","낮~중","B5"],
 ["6. 예방·위험요인","6.1 위험요인 묶음","결합","검진+가족력+생활+측정","대사증후군 등 통합 관리","중","A1·B4"],
 ["6. 예방·위험요인","6.2 암검진 갭(위험군)","결합","C형간염×간암검진","위험군 맞춤 검진 안내","중","A6"],
]
for i,r in enumerate(rows):
    ws.append(r)
    star = "★" in str(r[5]); combo = r[2].startswith("결합")
    for c in ws[ws.max_row]:
        c.font=Font(name=FONT,size=10); c.alignment=Alignment(vertical="center",wrap_text=True); c.border=B
        c.fill = STAR if star else (COMBO if combo else (ALT if i%2 else PatternFill()))
for col,w in zip("ABCDEFG",[18,24,12,24,28,12,14]): ws.column_dimensions[col].width=w
ws.freeze_panes="A2"
# 범례 시트
ws2=wb.create_sheet("범례")
for i,(k,v) in enumerate([("유형","단일=한 소스 / 결합=여러 소스(온톨로지 signalKey·교차)"),
   ("우선순위","★ 킬러 / 높음 / 중 / 낮음 (가치×난이도·리스크)"),
   ("안전(공리)","A=PHR, B=장치·환경, C=OCR, D=대화 — 통합 온톨로지 공리 참조"),
   ("색","주황=킬러(★) / 연두=결합형 서비스"),
   ("핵심","결합형(3·4·5·6층)이 차별점 — 범용 AI 불가, 온톨로지 있어야 안전 가능")],start=1):
    ws2[f"A{i}"]=k; ws2[f"B{i}"]=v
    ws2[f"A{i}"].font=Font(name=FONT,bold=True,size=10); ws2[f"B{i}"].font=Font(name=FONT,size=10)
ws2.column_dimensions["A"].width=14; ws2.column_dimensions["B"].width=80
wb.save("docs/서비스카탈로그_2026-06-21.xlsx"); print("saved", ws.max_row-1, "services")
